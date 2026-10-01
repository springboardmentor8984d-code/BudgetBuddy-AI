from django.db.models.functions import TruncMonth
from django.utils.dateparse import parse_date
from django.http import HttpResponse, JsonResponse
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from openpyxl import Workbook
from decimal import Decimal

from django.contrib.auth.models import User
from django.utils import timezone
from django.db.models import Sum

from rest_framework import generics, permissions, status
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework.permissions import AllowAny

from rest_framework_simplejwt.tokens import RefreshToken

from .models import (
    Income,
    Expense,
    Budget,
    SavingsGoal,
    Notification
)

from .serializers import (
    RegisterSerializer,
    IncomeSerializer,
    ExpenseSerializer,
    BudgetSerializer,
    SavingsGoalSerializer,
    NotificationSerializer
)


# ============================================================
# HOME
# ============================================================

def home(request):

    return JsonResponse({
        "message": "Welcome to BudgetBuddy API!"
    })


# ============================================================
# USER REGISTRATION
# ============================================================

class RegisterView(generics.CreateAPIView):

    queryset = User.objects.all()

    serializer_class = RegisterSerializer

    permission_classes = [
        AllowAny
    ]


# ============================================================
# USER LOGIN
# ============================================================

class LoginView(APIView):

    permission_classes = [
        AllowAny
    ]

    def post(self, request):

        username = request.data.get("username")

        password = request.data.get("password")

        try:

            user = User.objects.get(
                username=username
            )

        except User.DoesNotExist:

            return Response(
                {
                    "error":
                        "Invalid username or password"
                },
                status=status.HTTP_401_UNAUTHORIZED
            )

        if not user.check_password(password):

            return Response(
                {
                    "error":
                        "Invalid username or password"
                },
                status=status.HTTP_401_UNAUTHORIZED
            )

        refresh = RefreshToken.for_user(user)

        return Response({

            "message":
                "Login successful",

            "access":
                str(refresh.access_token),

            "refresh":
                str(refresh),

            "username":
                user.username,

            "email":
                user.email,

            "role":
                user.profile.role
        })


# ============================================================
# INCOME
# ============================================================

class IncomeListCreateView(
    generics.ListCreateAPIView
):

    serializer_class = IncomeSerializer

    permission_classes = [
        permissions.IsAuthenticated
    ]

    def get_queryset(self):

        return Income.objects.filter(
            user=self.request.user
        )

    def perform_create(self, serializer):

        serializer.save(
            user=self.request.user
        )


class IncomeDetailView(
    generics.RetrieveUpdateDestroyAPIView
):

    serializer_class = IncomeSerializer

    permission_classes = [
        permissions.IsAuthenticated
    ]

    def get_queryset(self):

        return Income.objects.filter(
            user=self.request.user
        )


# ============================================================
# EXPENSE
# ============================================================

class ExpenseListCreateView(
    generics.ListCreateAPIView
):

    serializer_class = ExpenseSerializer

    permission_classes = [
        permissions.IsAuthenticated
    ]

    def get_queryset(self):

        return Expense.objects.filter(
            user=self.request.user
        ).order_by(
            '-date',
            '-created_at'
        )

    def perform_create(self, serializer):

        expense = serializer.save(
            user=self.request.user
        )

        # ----------------------------------------------------
        # BUDGET ALERT CHECK
        # ----------------------------------------------------

        budgets = Budget.objects.filter(
            user=self.request.user,
            category=expense.category,
            month__year=expense.date.year,
            month__month=expense.date.month
        )

        for budget in budgets:

            create_budget_limit_notification(
                budget
            )


class ExpenseDetailView(
    generics.RetrieveUpdateDestroyAPIView
):

    serializer_class = ExpenseSerializer

    permission_classes = [
        permissions.IsAuthenticated
    ]

    def get_queryset(self):

        return Expense.objects.filter(
            user=self.request.user
        )


# ============================================================
# BUDGET
# ============================================================

class BudgetListCreateView(
    generics.ListCreateAPIView
):

    serializer_class = BudgetSerializer

    permission_classes = [
        permissions.IsAuthenticated
    ]

    def get_queryset(self):

        return Budget.objects.filter(
            user=self.request.user
        )

    def perform_create(self, serializer):

        serializer.save(
            user=self.request.user
        )


class BudgetDetailView(
    generics.RetrieveUpdateDestroyAPIView
):

    serializer_class = BudgetSerializer

    permission_classes = [
        permissions.IsAuthenticated
    ]

    def get_queryset(self):

        return Budget.objects.filter(
            user=self.request.user
        )


# ============================================================
# BUDGET UTILIZATION
# ============================================================

class BudgetUtilizationView(APIView):

    permission_classes = [
        permissions.IsAuthenticated
    ]

    def get(self, request):

        budgets = Budget.objects.filter(
            user=request.user
        )

        result = []

        for budget in budgets:

            spent = Expense.objects.filter(
                user=request.user,
                category=budget.category,
                date__year=budget.month.year,
                date__month=budget.month.month
            ).aggregate(
                total=Sum('amount')
            )['total'] or Decimal("0.00")

            remaining = (
                budget.amount -
                spent
            )

            if budget.amount > 0:

                utilization_percentage = (
                    spent /
                    budget.amount
                ) * Decimal("100")

            else:

                utilization_percentage = Decimal("0.00")

            if spent > budget.amount:

                budget_status = "Overspending"

            elif utilization_percentage >= Decimal("80"):

                budget_status = "Approaching Limit"

            else:

                budget_status = "Within Budget"

            result.append({

                "category":
                    budget.category,

                "month":
                    budget.month,

                "budget_amount":
                    budget.amount,

                "spent_amount":
                    spent,

                "remaining_amount":
                    remaining,

                "utilization_percentage":
                    utilization_percentage,

                "status":
                    budget_status
            })

        return Response(result)


# ============================================================
# TRANSACTION DASHBOARD
# ============================================================

class TransactionDashboardView(APIView):

    permission_classes = [
        permissions.IsAuthenticated
    ]

    def get(self, request):

        incomes = Income.objects.filter(
            user=request.user
        )

        expenses = Expense.objects.filter(
            user=request.user
        )

        total_income = incomes.aggregate(
            total=Sum('amount')
        )['total'] or Decimal("0.00")

        total_expenses = expenses.aggregate(
            total=Sum('amount')
        )['total'] or Decimal("0.00")

        remaining_amount = (
            total_income -
            total_expenses
        )

        transactions = []

        for income in incomes:

            transactions.append({

                "type":
                    "Income",

                "id":
                    income.id,

                "source":
                    income.source,

                "amount":
                    income.amount,

                "date":
                    income.date,

                "description":
                    income.description
            })

        for expense in expenses:

            transactions.append({

                "type":
                    "Expense",

                "id":
                    expense.id,

                "category":
                    expense.category,

                "amount":
                    expense.amount,

                "date":
                    expense.date,

                "description":
                    expense.description
            })

        transactions.sort(
            key=lambda transaction:
                transaction["date"],
            reverse=True
        )

        return Response({

            "total_income":
                total_income,

            "total_expenses":
                total_expenses,

            "remaining_amount":
                remaining_amount,

            "recent_transactions":
                transactions[:10]
        })


# ============================================================
# ANALYTICS
# ============================================================

class AnalyticsView(APIView):

    permission_classes = [
        permissions.IsAuthenticated
    ]

    def get(self, request):

        start_date = request.query_params.get(
            'start_date'
        )

        end_date = request.query_params.get(
            'end_date'
        )

        month = request.query_params.get(
            'month'
        )

        # ----------------------------------------------------
        # VALIDATE START DATE
        # ----------------------------------------------------

        if start_date:

            start_date = parse_date(
                start_date
            )

            if start_date is None:

                return Response(
                    {
                        "detail":
                            "Invalid start_date. "
                            "Use YYYY-MM-DD."
                    },
                    status=status.HTTP_400_BAD_REQUEST
                )

        # ----------------------------------------------------
        # VALIDATE END DATE
        # ----------------------------------------------------

        if end_date:

            end_date = parse_date(
                end_date
            )

            if end_date is None:

                return Response(
                    {
                        "detail":
                            "Invalid end_date. "
                            "Use YYYY-MM-DD."
                    },
                    status=status.HTTP_400_BAD_REQUEST
                )

        # ----------------------------------------------------
        # VALIDATE DATE RANGE
        # ----------------------------------------------------

        if (
            start_date
            and end_date
            and start_date > end_date
        ):

            return Response(
                {
                    "detail":
                        "start_date cannot be "
                        "after end_date."
                },
                status=status.HTTP_400_BAD_REQUEST
            )

        # ----------------------------------------------------
        # GET USER DATA
        # ----------------------------------------------------

        incomes = Income.objects.filter(
            user=request.user
        )

        expenses = Expense.objects.filter(
            user=request.user
        )

        # ----------------------------------------------------
        # DATE FILTERING
        # ----------------------------------------------------

        if start_date:

            incomes = incomes.filter(
                date__gte=start_date
            )

            expenses = expenses.filter(
                date__gte=start_date
            )

        if end_date:

            incomes = incomes.filter(
                date__lte=end_date
            )

            expenses = expenses.filter(
                date__lte=end_date
            )

        # ----------------------------------------------------
        # MONTH FILTER
        # ----------------------------------------------------

        if month:

            try:

                year, month_number = (
                    month.split('-')
                )

                year = int(year)

                month_number = int(
                    month_number
                )

                if (
                    len(str(year)) != 4
                    or month_number < 1
                    or month_number > 12
                ):

                    raise ValueError

            except (ValueError, TypeError):

                return Response(
                    {
                        "detail":
                            "Invalid month. "
                            "Use YYYY-MM."
                    },
                    status=status.HTTP_400_BAD_REQUEST
                )

            incomes = incomes.filter(
                date__year=year,
                date__month=month_number
            )

            expenses = expenses.filter(
                date__year=year,
                date__month=month_number
            )

        # ----------------------------------------------------
        # TOTAL INCOME
        # ----------------------------------------------------

        total_income = incomes.aggregate(
            total=Sum('amount')
        )['total'] or Decimal("0.00")

        # ----------------------------------------------------
        # TOTAL EXPENSES
        # ----------------------------------------------------

        total_expenses = expenses.aggregate(
            total=Sum('amount')
        )['total'] or Decimal("0.00")

        # ----------------------------------------------------
        # REMAINING AMOUNT
        # ----------------------------------------------------

        remaining_amount = (
            total_income -
            total_expenses
        )

        # ----------------------------------------------------
        # CATEGORY SUMMARY
        # ----------------------------------------------------

        category_summary = []

        category_data = (
            expenses
            .values('category')
            .annotate(
                total=Sum('amount')
            )
            .order_by('category')
        )

        for item in category_data:

            category_summary.append({

                "category":
                    item['category'],

                "total":
                    item['total']
                    or Decimal("0.00")
            })

        # ----------------------------------------------------
        # MONTHLY TRENDS
        # ----------------------------------------------------

        income_monthly = (
            incomes
            .annotate(
                month_value=TruncMonth(
                    'date'
                )
            )
            .values('month_value')
            .annotate(
                total=Sum('amount')
            )
            .order_by('month_value')
        )

        expense_monthly = (
            expenses
            .annotate(
                month_value=TruncMonth(
                    'date'
                )
            )
            .values('month_value')
            .annotate(
                total=Sum('amount')
            )
            .order_by('month_value')
        )

        monthly_data = {}

        for item in income_monthly:

            month_key = item[
                'month_value'
            ].strftime('%Y-%m')

            if month_key not in monthly_data:

                monthly_data[month_key] = {

                    "month":
                        month_key,

                    "income":
                        Decimal("0.00"),

                    "expenses":
                        Decimal("0.00"),

                    "savings":
                        Decimal("0.00")
                }

            monthly_data[
                month_key
            ]["income"] = (
                item['total']
                or Decimal("0.00")
            )

        for item in expense_monthly:

            month_key = item[
                'month_value'
            ].strftime('%Y-%m')

            if month_key not in monthly_data:

                monthly_data[month_key] = {

                    "month":
                        month_key,

                    "income":
                        Decimal("0.00"),

                    "expenses":
                        Decimal("0.00"),

                    "savings":
                        Decimal("0.00")
                }

            monthly_data[
                month_key
            ]["expenses"] = (
                item['total']
                or Decimal("0.00")
            )

        for item in monthly_data.values():

            item["savings"] = (
                item["income"] -
                item["expenses"]
            )

        monthly_trends = sorted(
            monthly_data.values(),
            key=lambda item:
                item["month"]
        )

        return Response({

            "filters": {

                "start_date":
                    start_date,

                "end_date":
                    end_date,

                "month":
                    month
            },

            "summary": {

                "total_income":
                    total_income,

                "total_expenses":
                    total_expenses,

                "remaining_amount":
                    remaining_amount
            },

            "category_wise_expenses":
                category_summary,

            "monthly_trends":
                monthly_trends
        })


# ============================================================
# SAVINGS GOAL
# ============================================================

class SavingsGoalListCreateView(
    generics.ListCreateAPIView
):

    serializer_class = SavingsGoalSerializer

    permission_classes = [
        permissions.IsAuthenticated
    ]

    def get_queryset(self):

        return SavingsGoal.objects.filter(
            user=self.request.user
        ).order_by('-id')

    def perform_create(self, serializer):

        goal = serializer.save(
            user=self.request.user
        )

        create_savings_milestone_notification(
            goal
        )


class SavingsGoalDetailView(
    generics.RetrieveUpdateDestroyAPIView
):

    serializer_class = SavingsGoalSerializer

    permission_classes = [
        permissions.IsAuthenticated
    ]

    def get_queryset(self):

        return SavingsGoal.objects.filter(
            user=self.request.user
        )

    def perform_update(self, serializer):

        goal = serializer.save()

        create_savings_milestone_notification(
            goal
        )


# ============================================================
# SAVINGS GOAL PROGRESS
# ============================================================

class SavingsGoalProgressView(APIView):

    permission_classes = [
        permissions.IsAuthenticated
    ]

    def patch(self, request, pk):

        try:

            goal = SavingsGoal.objects.get(
                pk=pk,
                user=request.user
            )

        except SavingsGoal.DoesNotExist:

            return Response(
                {
                    "detail":
                        "Savings goal not found."
                },
                status=status.HTTP_404_NOT_FOUND
            )

        current_amount = request.data.get(
            "current_amount"
        )

        if current_amount is None:

            return Response(
                {
                    "current_amount":
                        "This field is required."
                },
                status=status.HTTP_400_BAD_REQUEST
            )

        serializer = SavingsGoalSerializer(
            goal,
            data={
                "current_amount":
                    current_amount
            },
            partial=True,
            context={
                "request":
                    request
            }
        )

        if not serializer.is_valid():

            return Response(
                serializer.errors,
                status=status.HTTP_400_BAD_REQUEST
            )

        goal = serializer.save()

        create_savings_milestone_notification(
            goal
        )

        return Response(
            serializer.data,
            status=status.HTTP_200_OK
        )


# ============================================================
# SAVINGS GOAL COMPLETION
# ============================================================

class SavingsGoalCompletionView(APIView):

    permission_classes = [
        permissions.IsAuthenticated
    ]

    def get(self, request, pk):

        try:

            goal = SavingsGoal.objects.get(
                pk=pk,
                user=request.user
            )

        except SavingsGoal.DoesNotExist:

            return Response(
                {
                    "detail":
                        "Savings goal not found."
                },
                status=status.HTTP_404_NOT_FOUND
            )

        return Response({

            "id":
                goal.id,

            "goal_name":
                goal.goal_name,

            "target_amount":
                goal.target_amount,

            "current_amount":
                goal.current_amount,

            "progress_percentage":
                goal.progress_percentage,

            "status":
                goal.status,

            "completed":
                goal.status == "completed"
        })


# ============================================================
# NOTIFICATION
# ============================================================

class NotificationListCreateView(
    generics.ListCreateAPIView
):

    serializer_class = NotificationSerializer

    permission_classes = [
        permissions.IsAuthenticated
    ]

    def get_queryset(self):

        return Notification.objects.filter(
            user=self.request.user
        ).order_by(
            '-created_at'
        )

    def perform_create(self, serializer):

        serializer.save(
            user=self.request.user
        )


# ============================================================
# VIEW UNREAD NOTIFICATIONS
# ============================================================

class UnreadNotificationView(
    generics.ListAPIView
):

    serializer_class = NotificationSerializer

    permission_classes = [
        permissions.IsAuthenticated
    ]

    def get_queryset(self):

        return Notification.objects.filter(
            user=self.request.user,
            is_read=False
        ).order_by(
            '-created_at'
        )


# ============================================================
# MARK NOTIFICATION AS READ
# ============================================================

class NotificationReadView(APIView):

    permission_classes = [
        permissions.IsAuthenticated
    ]

    def patch(self, request, pk):

        try:

            notification = Notification.objects.get(
                pk=pk,
                user=request.user
            )

        except Notification.DoesNotExist:

            return Response(
                {
                    "detail":
                        "Notification not found."
                },
                status=status.HTTP_404_NOT_FOUND
            )

        notification.is_read = True

        notification.save(
            update_fields=[
                'is_read'
            ]
        )

        return Response(
            NotificationSerializer(
                notification
            ).data,
            status=status.HTTP_200_OK
        )


# ============================================================
# BUDGET ALERT TRIGGER
# ============================================================

def create_budget_limit_notification(budget):

    """
    Checks budget utilization.

    Below 80%:
        No notification

    80%-100%:
        Approaching Limit

    Above 100%:
        Overspending

    Duplicate notifications are prevented.
    """

    spent = Expense.objects.filter(
        user=budget.user,
        category=budget.category,
        date__year=budget.month.year,
        date__month=budget.month.month
    ).aggregate(
        total=Sum('amount')
    )['total'] or Decimal("0.00")

    if budget.amount > 0:

        utilization_percentage = (
            spent /
            budget.amount
        ) * Decimal("100")

    else:

        utilization_percentage = Decimal("0.00")

    if spent > budget.amount:

        alert_status = "Overspending"

    elif utilization_percentage >= Decimal("80"):

        alert_status = "Approaching Limit"

    else:

        alert_status = "Within Budget"

    if utilization_percentage < Decimal("80"):

        return

    already_exists = Notification.objects.filter(
        user=budget.user,
        notification_type='budget_limit',
        related_object_id=budget.id
    ).exists()

    if already_exists:

        return

    if alert_status == "Overspending":

        message = (
            f"Your {budget.category} budget has "
            f"been exceeded. You have spent "
            f"₹{spent} out of ₹{budget.amount}."
        )

    else:

        message = (
            f"Your {budget.category} budget is "
            f"approaching its limit. You have "
            f"spent ₹{spent} out of "
            f"₹{budget.amount}."
        )

    Notification.objects.create(

        user=budget.user,

        notification_type=
            'budget_limit',

        message=message,

        related_object_id=
            budget.id
    )


# ============================================================
# SAVINGS MILESTONE NOTIFICATION
# ============================================================

def create_savings_milestone_notification(goal):

    """
    Creates savings milestone notifications
    at 50% and 100%.

    Duplicate notifications are prevented.
    """

    progress = goal.progress_percentage

    milestone = None

    if progress >= Decimal("100"):

        milestone = "100%"

    elif progress >= Decimal("50"):

        milestone = "50%"

    if milestone is None:

        return

    already_exists = Notification.objects.filter(
        user=goal.user,
        notification_type='goal_milestone',
        related_object_id=goal.id,
        message__contains=milestone
    ).exists()

    if already_exists:

        return

    Notification.objects.create(

        user=goal.user,

        notification_type=
            'goal_milestone',

        message=(
            f"Your savings goal "
            f"'{goal.goal_name}' "
            f"has reached {milestone}."
        ),

        related_object_id=
            goal.id
    )


# ============================================================
# MONTHLY REPORT NOTIFICATION
# ============================================================

def create_monthly_report_notification(
    user,
    month
):

    """
    Creates a notification when the user's
    monthly financial report is generated.

    Only one monthly report notification is
    created for each user and month.
    """

    already_exists = Notification.objects.filter(
        user=user,
        notification_type='monthly_report',
        message__contains=month
    ).exists()

    if already_exists:

        return

    Notification.objects.create(

        user=user,

        notification_type=
            'monthly_report',

        message=(
            f"Your monthly financial report "
            f"for {month} is ready."
        )
    )


# ============================================================
# MONTHLY FINANCIAL REPORT
# ============================================================

class MonthlyReportView(APIView):

    permission_classes = [
        permissions.IsAuthenticated
    ]

    def get(self, request):

        # ----------------------------------------------------
        # GET MONTH
        # ----------------------------------------------------

        month = request.query_params.get(
            'month'
        )

        # ----------------------------------------------------
        # DEFAULT CURRENT MONTH
        # ----------------------------------------------------

        if not month:

            today = timezone.localdate()

            year = today.year

            month_number = today.month

            month = (
                f"{year:04d}-{month_number:02d}"
            )

        # ----------------------------------------------------
        # VALIDATE MONTH
        # ----------------------------------------------------

        try:

            year, month_number = (
                month.split('-')
            )

            year = int(year)

            month_number = int(
                month_number
            )

            if (
                len(str(year)) != 4
                or month_number < 1
                or month_number > 12
            ):

                raise ValueError

        except (ValueError, TypeError):

            return Response(
                {
                    "detail":
                        "Invalid month. "
                        "Use YYYY-MM."
                },
                status=status.HTTP_400_BAD_REQUEST
            )

        # ----------------------------------------------------
        # GET USER DATA
        # ----------------------------------------------------

        incomes = Income.objects.filter(
            user=request.user,
            date__year=year,
            date__month=month_number
        )

        expenses = Expense.objects.filter(
            user=request.user,
            date__year=year,
            date__month=month_number
        )

        budgets = Budget.objects.filter(
            user=request.user,
            month__year=year,
            month__month=month_number
        )

        savings_goals = SavingsGoal.objects.filter(
            user=request.user
        )

        # ----------------------------------------------------
        # TOTAL INCOME
        # ----------------------------------------------------

        total_income = incomes.aggregate(
            total=Sum('amount')
        )['total'] or Decimal("0.00")

        # ----------------------------------------------------
        # TOTAL EXPENSES
        # ----------------------------------------------------

        total_expenses = expenses.aggregate(
            total=Sum('amount')
        )['total'] or Decimal("0.00")

        # ----------------------------------------------------
        # SAVINGS
        # ----------------------------------------------------

        savings = (
            total_income -
            total_expenses
        )

        # ----------------------------------------------------
        # CATEGORY-WISE EXPENSES
        # ----------------------------------------------------

        category_data = (
            expenses
            .values('category')
            .annotate(
                total=Sum('amount')
            )
            .order_by('category')
        )

        category_wise_expenses = []

        for item in category_data:

            category_wise_expenses.append({

                "category":
                    item['category'],

                "total":
                    item['total']
                    or Decimal("0.00")
            })

        # ----------------------------------------------------
        # BUDGET SUMMARY
        # ----------------------------------------------------

        budget_summary = []

        for budget in budgets:

            spent = expenses.filter(
                category=budget.category
            ).aggregate(
                total=Sum('amount')
            )['total'] or Decimal("0.00")

            remaining = (
                budget.amount -
                spent
            )

            if spent > budget.amount:

                budget_status = "Overspending"

            elif spent >= (
                budget.amount *
                Decimal("0.80")
            ):

                budget_status = "Near Limit"

            else:

                budget_status = "Within Budget"

            budget_summary.append({

                "category":
                    budget.category,

                "budget_amount":
                    budget.amount,

                "spent_amount":
                    spent,

                "remaining_amount":
                    remaining,

                "status":
                    budget_status
            })

        # ----------------------------------------------------
        # SAVINGS GOALS
        # ----------------------------------------------------

        savings_goal_summary = []

        for goal in savings_goals:

            savings_goal_summary.append({

                "id":
                    goal.id,

                "goal_name":
                    goal.goal_name,

                "target_amount":
                    goal.target_amount,

                "current_amount":
                    goal.current_amount,

                "progress_percentage":
                    goal.progress_percentage,

                "status":
                    goal.status,

                "deadline":
                    goal.deadline
            })

        # ----------------------------------------------------
        # CREATE MONTHLY REPORT NOTIFICATION
        # ----------------------------------------------------

        create_monthly_report_notification(
            request.user,
            month
        )

        # ----------------------------------------------------
        # RESPONSE
        # ----------------------------------------------------

        return Response({

            "report": {

                "report_type":
                    "Monthly Financial Report",

                "month":
                    month
            },

            "summary": {

                "total_income":
                    total_income,

                "total_expenses":
                    total_expenses,

                "savings":
                    savings
            },

            "category_wise_expenses":
                category_wise_expenses,

            "budget_summary":
                budget_summary,

            "savings_goals":
                savings_goal_summary
        })


# ============================================================
# MONTHLY FINANCIAL REPORT PDF
# ============================================================

class MonthlyReportPDFView(APIView):

    permission_classes = [
        permissions.IsAuthenticated
    ]

    def get(self, request):

        # ----------------------------------------------------
        # GET MONTH
        # ----------------------------------------------------

        month = request.query_params.get(
            'month'
        )

        # ----------------------------------------------------
        # DEFAULT CURRENT MONTH
        # ----------------------------------------------------

        if not month:

            today = timezone.localdate()

            year = today.year

            month_number = today.month

            month = (
                f"{year:04d}-{month_number:02d}"
            )

        # ----------------------------------------------------
        # VALIDATE MONTH
        # ----------------------------------------------------

        try:

            year, month_number = (
                month.split('-')
            )

            year = int(year)

            month_number = int(
                month_number
            )

            if (
                len(str(year)) != 4
                or month_number < 1
                or month_number > 12
            ):

                raise ValueError

        except (ValueError, TypeError):

            return Response(
                {
                    "detail":
                        "Invalid month. "
                        "Use YYYY-MM."
                },
                status=status.HTTP_400_BAD_REQUEST
            )

        # ----------------------------------------------------
        # GET USER DATA
        # ----------------------------------------------------

        incomes = Income.objects.filter(
            user=request.user,
            date__year=year,
            date__month=month_number
        )

        expenses = Expense.objects.filter(
            user=request.user,
            date__year=year,
            date__month=month_number
        )

        budgets = Budget.objects.filter(
            user=request.user,
            month__year=year,
            month__month=month_number
        )

        savings_goals = SavingsGoal.objects.filter(
            user=request.user
        )

        # ----------------------------------------------------
        # TOTALS
        # ----------------------------------------------------

        total_income = incomes.aggregate(
            total=Sum('amount')
        )['total'] or Decimal("0.00")

        total_expenses = expenses.aggregate(
            total=Sum('amount')
        )['total'] or Decimal("0.00")

        savings = (
            total_income -
            total_expenses
        )

        # ----------------------------------------------------
        # CATEGORY DATA
        # ----------------------------------------------------

        category_data = (
            expenses
            .values('category')
            .annotate(
                total=Sum('amount')
            )
            .order_by('category')
        )

        # ----------------------------------------------------
        # CREATE PDF RESPONSE
        # ----------------------------------------------------

        response = HttpResponse(
            content_type='application/pdf'
        )

        response[
            'Content-Disposition'
        ] = (
            'attachment; '
            f'filename="monthly_report_{month}.pdf"'
        )

        # ----------------------------------------------------
        # CREATE PDF
        # ----------------------------------------------------

        pdf = canvas.Canvas(
            response,
            pagesize=A4
        )

        width, height = A4

        y = height - 50

        # ----------------------------------------------------
        # TITLE
        # ----------------------------------------------------

        pdf.setFont(
            "Helvetica-Bold",
            18
        )

        pdf.drawString(
            50,
            y,
            "BudgetBuddy"
        )

        y -= 30

        pdf.setFont(
            "Helvetica-Bold",
            14
        )

        pdf.drawString(
            50,
            y,
            "Monthly Financial Report"
        )

        y -= 25

        pdf.setFont(
            "Helvetica",
            11
        )

        pdf.drawString(
            50,
            y,
            f"Month: {month}"
        )

        y -= 35

        # ----------------------------------------------------
        # SUMMARY
        # ----------------------------------------------------

        pdf.setFont(
            "Helvetica-Bold",
            13
        )

        pdf.drawString(
            50,
            y,
            "Financial Summary"
        )

        y -= 25

        pdf.setFont(
            "Helvetica",
            11
        )

        pdf.drawString(
            60,
            y,
            f"Total Income: Rs. {total_income}"
        )

        y -= 20

        pdf.drawString(
            60,
            y,
            f"Total Expenses: Rs. {total_expenses}"
        )

        y -= 20

        pdf.drawString(
            60,
            y,
            f"Savings: Rs. {savings}"
        )

        y -= 35

        # ----------------------------------------------------
        # CATEGORY-WISE EXPENSES
        # ----------------------------------------------------

        pdf.setFont(
            "Helvetica-Bold",
            13
        )

        pdf.drawString(
            50,
            y,
            "Category-wise Expenses"
        )

        y -= 25

        pdf.setFont(
            "Helvetica",
            11
        )

        for item in category_data:

            category = item[
                'category'
            ]

            total = (
                item['total']
                or Decimal("0.00")
            )

            pdf.drawString(
                60,
                y,
                f"{category}: Rs. {total}"
            )

            y -= 20

            if y < 60:

                pdf.showPage()

                y = height - 50

                pdf.setFont(
                    "Helvetica",
                    11
                )

        # ----------------------------------------------------
        # BUDGET SUMMARY
        # ----------------------------------------------------

        y -= 15

        pdf.setFont(
            "Helvetica-Bold",
            13
        )

        pdf.drawString(
            50,
            y,
            "Budget Summary"
        )

        y -= 25

        pdf.setFont(
            "Helvetica",
            11
        )

        for budget in budgets:

            spent = expenses.filter(
                category=budget.category
            ).aggregate(
                total=Sum('amount')
            )['total'] or Decimal("0.00")

            remaining = (
                budget.amount -
                spent
            )

            if spent > budget.amount:

                budget_status = "Overspending"

            elif spent >= (
                budget.amount *
                Decimal("0.80")
            ):

                budget_status = "Near Limit"

            else:

                budget_status = "Within Budget"

            pdf.drawString(
                60,
                y,
                f"{budget.category}: "
                f"Budget Rs. {budget.amount}, "
                f"Spent Rs. {spent}, "
                f"Remaining Rs. {remaining}"
            )

            y -= 20

            pdf.drawString(
                60,
                y,
                f"Status: {budget_status}"
            )

            y -= 25

            if y < 60:

                pdf.showPage()

                y = height - 50

                pdf.setFont(
                    "Helvetica",
                    11
                )

        # ----------------------------------------------------
        # SAVINGS GOALS
        # ----------------------------------------------------

        y -= 10

        pdf.setFont(
            "Helvetica-Bold",
            13
        )

        pdf.drawString(
            50,
            y,
            "Savings Goals"
        )

        y -= 25

        pdf.setFont(
            "Helvetica",
            11
        )

        for goal in savings_goals:

            pdf.drawString(
                60,
                y,
                f"{goal.goal_name}: "
                f"Rs. {goal.current_amount} / "
                f"Rs. {goal.target_amount}"
            )

            y -= 20

            pdf.drawString(
                60,
                y,
                f"Progress: "
                f"{goal.progress_percentage}% "
                f"| Status: {goal.status}"
            )

            y -= 25

            if y < 60:

                pdf.showPage()

                y = height - 50

                pdf.setFont(
                    "Helvetica",
                    11
                )

        # ----------------------------------------------------
        # FOOTER
        # ----------------------------------------------------

        pdf.setFont(
            "Helvetica",
            9
        )

        pdf.setFillColor(
            colors.grey
        )

        pdf.drawString(
            50,
            30,
            "Generated by BudgetBuddy"
        )

        # ----------------------------------------------------
        # SAVE PDF
        # ----------------------------------------------------

        pdf.save()

        # ----------------------------------------------------
        # CREATE MONTHLY REPORT NOTIFICATION
        # ----------------------------------------------------

        create_monthly_report_notification(
            request.user,
            month
        )

        return response


# ============================================================
# MONTHLY FINANCIAL REPORT - EXCEL
# ============================================================

class MonthlyReportExcelView(APIView):

    permission_classes = [
        permissions.IsAuthenticated
    ]

    def get(self, request):

        # ----------------------------------------------------
        # GET MONTH
        # ----------------------------------------------------

        month = request.query_params.get(
            'month'
        )

        # ----------------------------------------------------
        # DEFAULT CURRENT MONTH
        # ----------------------------------------------------

        if not month:

            today = timezone.localdate()

            year = today.year

            month_number = today.month

            month = (
                f"{year:04d}-{month_number:02d}"
            )

        # ----------------------------------------------------
        # VALIDATE MONTH
        # ----------------------------------------------------

        try:

            year, month_number = (
                month.split('-')
            )

            year = int(year)

            month_number = int(
                month_number
            )

            if (
                len(str(year)) != 4
                or month_number < 1
                or month_number > 12
            ):

                raise ValueError

        except (ValueError, TypeError):

            return Response(
                {
                    "detail":
                        "Invalid month. "
                        "Use YYYY-MM."
                },
                status=status.HTTP_400_BAD_REQUEST
            )

        # ----------------------------------------------------
        # GET USER DATA
        # ----------------------------------------------------

        incomes = Income.objects.filter(
            user=request.user,
            date__year=year,
            date__month=month_number
        )

        expenses = Expense.objects.filter(
            user=request.user,
            date__year=year,
            date__month=month_number
        )

        budgets = Budget.objects.filter(
            user=request.user,
            month__year=year,
            month__month=month_number
        )

        savings_goals = SavingsGoal.objects.filter(
            user=request.user
        )

        # ----------------------------------------------------
        # CALCULATE SUMMARY
        # ----------------------------------------------------

        total_income = incomes.aggregate(
            total=Sum('amount')
        )['total'] or Decimal("0.00")

        total_expenses = expenses.aggregate(
            total=Sum('amount')
        )['total'] or Decimal("0.00")

        savings = (
            total_income -
            total_expenses
        )

        # ----------------------------------------------------
        # CREATE WORKBOOK
        # ----------------------------------------------------

        workbook = Workbook()

        sheet = workbook.active

        sheet.title = "Monthly Report"

        # ----------------------------------------------------
        # TITLE
        # ----------------------------------------------------

        sheet["A1"] = (
            "BudgetBuddy Monthly Financial Report"
        )

        sheet["A2"] = (
            f"Month: {month}"
        )

        # ----------------------------------------------------
        # FINANCIAL SUMMARY
        # ----------------------------------------------------

        sheet["A4"] = "Financial Summary"

        sheet["A5"] = "Total Income"

        sheet["B5"] = float(
            total_income
        )

        sheet["A6"] = "Total Expenses"

        sheet["B6"] = float(
            total_expenses
        )

        sheet["A7"] = "Savings"

        sheet["B7"] = float(
            savings
        )

        # ----------------------------------------------------
        # CATEGORY-WISE EXPENSES
        # ----------------------------------------------------

        sheet["A9"] = (
            "Category-wise Expenses"
        )

        sheet["A10"] = "Category"

        sheet["B10"] = "Amount"

        row = 11

        category_data = (
            expenses
            .values("category")
            .annotate(
                total=Sum("amount")
            )
            .order_by("category")
        )

        for item in category_data:

            sheet.cell(
                row=row,
                column=1,
                value=item["category"]
            )

            sheet.cell(
                row=row,
                column=2,
                value=float(
                    item["total"]
                    or Decimal("0.00")
                )
            )

            row += 1

        # ----------------------------------------------------
        # BUDGET SUMMARY
        # ----------------------------------------------------

        row += 1

        sheet.cell(
            row=row,
            column=1,
            value="Budget Summary"
        )

        row += 1

        sheet.cell(
            row=row,
            column=1,
            value="Category"
        )

        sheet.cell(
            row=row,
            column=2,
            value="Budget"
        )

        sheet.cell(
            row=row,
            column=3,
            value="Spent"
        )

        sheet.cell(
            row=row,
            column=4,
            value="Remaining"
        )

        sheet.cell(
            row=row,
            column=5,
            value="Status"
        )

        row += 1

        for budget in budgets:

            spent = expenses.filter(
                category=budget.category
            ).aggregate(
                total=Sum("amount")
            )["total"] or Decimal("0.00")

            remaining = (
                budget.amount -
                spent
            )

            if spent > budget.amount:

                budget_status = "Overspending"

            elif spent >= (
                budget.amount *
                Decimal("0.80")
            ):

                budget_status = "Near Limit"

            else:

                budget_status = "Within Budget"

            sheet.cell(
                row=row,
                column=1,
                value=budget.category
            )

            sheet.cell(
                row=row,
                column=2,
                value=float(
                    budget.amount
                )
            )

            sheet.cell(
                row=row,
                column=3,
                value=float(
                    spent
                )
            )

            sheet.cell(
                row=row,
                column=4,
                value=float(
                    remaining
                )
            )

            sheet.cell(
                row=row,
                column=5,
                value=budget_status
            )

            row += 1

        # ----------------------------------------------------
        # SAVINGS GOALS
        # ----------------------------------------------------

        row += 1

        sheet.cell(
            row=row,
            column=1,
            value="Savings Goals"
        )

        row += 1

        sheet.cell(
            row=row,
            column=1,
            value="Goal Name"
        )

        sheet.cell(
            row=row,
            column=2,
            value="Target Amount"
        )

        sheet.cell(
            row=row,
            column=3,
            value="Current Amount"
        )

        sheet.cell(
            row=row,
            column=4,
            value="Progress %"
        )

        sheet.cell(
            row=row,
            column=5,
            value="Status"
        )

        row += 1

        for goal in savings_goals:

            sheet.cell(
                row=row,
                column=1,
                value=goal.goal_name
            )

            sheet.cell(
                row=row,
                column=2,
                value=float(
                    goal.target_amount
                )
            )

            sheet.cell(
                row=row,
                column=3,
                value=float(
                    goal.current_amount
                )
            )

            sheet.cell(
                row=row,
                column=4,
                value=float(
                    goal.progress_percentage
                )
            )

            sheet.cell(
                row=row,
                column=5,
                value=goal.status
            )

            row += 1

        # ----------------------------------------------------
        # COLUMN WIDTHS
        # ----------------------------------------------------

        sheet.column_dimensions[
            "A"
        ].width = 30

        sheet.column_dimensions[
            "B"
        ].width = 18

        sheet.column_dimensions[
            "C"
        ].width = 18

        sheet.column_dimensions[
            "D"
        ].width = 18

        sheet.column_dimensions[
            "E"
        ].width = 20

        # ----------------------------------------------------
        # CREATE EXCEL RESPONSE
        # ----------------------------------------------------

        response = HttpResponse(
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )

        response[
            "Content-Disposition"
        ] = (
            'attachment; '
            f'filename="monthly_report_{month}.xlsx"'
        )

        workbook.save(
            response
        )

        # ----------------------------------------------------
        # CREATE MONTHLY REPORT NOTIFICATION
        # ----------------------------------------------------

        create_monthly_report_notification(
            request.user,
            month
        )

        return response